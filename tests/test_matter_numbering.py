"""The last safety check: matters numbered 1..N should come to N.

Not proof of anything on its own - firms delete matters without renumbering -
but a gap is also what an undercount looks like, so it is always reported.
"""

import pytest
from docx import Document

from conftest import SAMPLES
from parsing import parse_document


def submission(tmp_path, labels, name="doc.docx"):
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"

    for label in labels:
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = label
        table.rows[1].cells[0].text = "Name of client"
        table.rows[1].cells[1].text = "A client"
    path = tmp_path / name
    document.save(path)
    return parse_document(path).metrics["matters"]


def test_consecutive_numbering_raises_nothing(tmp_path):
    metric = submission(tmp_path, [f"Publishable matter {n}" for n in (1, 2, 3)])
    assert metric.value == 3
    assert metric.numbering_gap is False
    assert not any("numbered up to" in n for n in metric.notes)


def test_a_gap_is_reported(tmp_path):
    """The reported shape: 12 publishable matters but only 11 counted."""
    labels = [f"Publishable matter {n}" for n in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12)]
    metric = submission(tmp_path, labels)
    assert metric.value == 11
    assert metric.numbering_gap is True
    note = next(n for n in metric.notes if "numbered up to" in n)
    assert "numbered up to 12 but 11 were counted" in note
    assert "no 7" in note


def test_the_two_categories_are_checked_separately(tmp_path):
    labels = [f"Publishable matter {n}" for n in (1, 2)]
    labels += [f"Non-publishable matter {n}" for n in (1, 2, 4)]
    metric = submission(tmp_path, labels)
    assert metric.value == 5
    assert metric.numbering_gap is True
    notes = " ".join(metric.notes)
    assert "Non-publishable matters are numbered up to 4" in notes
    assert "Publishable matters are numbered up to" not in notes


def test_numbering_starting_at_two_is_reported(tmp_path):
    metric = submission(tmp_path, [f"Publishable matter {n}" for n in (2, 3, 4)])
    assert metric.numbering_gap is True
    assert "no 1" in " ".join(metric.notes)


def test_a_repeated_number_is_reported(tmp_path):
    metric = submission(tmp_path, ["Publishable matter 1", "Publishable matter 2",
                                   "Publishable matter 2"])
    assert metric.value == 3
    assert metric.numbering_gap is True
    assert "repeated 2" in " ".join(metric.notes)


def test_the_count_is_never_blanked_by_a_gap(tmp_path):
    """A gap is a prompt to look, not a reason to withhold the number."""
    metric = submission(tmp_path, [f"Publishable matter {n}" for n in (1, 3)])
    assert metric.value == 2
    assert not metric.is_unparsed


def test_an_empty_template_is_reported_by_its_own_note_not_as_a_gap(tmp_path):
    """Matter 3 is present but empty.

    The numbering check runs over the matters actually counted, so 1 and 2 are
    consecutive and no gap is raised - the skipped template already has its own
    note, and reporting it twice would be noise.
    """
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    for n in (1, 2, 3):
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = f"Publishable matter {n}"
        if n != 3:
            table.rows[1].cells[0].text = "Name of client"
            table.rows[1].cells[1].text = "A client"
    path = tmp_path / "empty.docx"
    document.save(path)

    metric = parse_document(path).metrics["matters"]
    assert metric.value == 2
    assert any("empty matter template" in n for n in metric.notes)
    assert metric.numbering_gap is False


@pytest.mark.parametrize("doc_key", ["alpha", "beta"])
def test_the_samples_have_clean_numbering(doc_key):
    """alpha and beta number their matters consecutively; gamma does not."""
    metric = parse_document(SAMPLES[doc_key]).metrics["matters"]
    assert metric.numbering_gap is False


def test_gamma_is_the_one_with_a_gap():
    metric = parse_document(SAMPLES["gamma"]).metrics["matters"]
    assert metric.numbering_gap is True


# --- templates that do not number the matter boxes -------------------------


def test_unnumbered_matter_labels_are_counted(tmp_path):
    """One firm's template labels every box "Publishable matter", no number.

    Requiring a number meant no matter table was found at all, so every
    document from that template reported the metric as uncountable and every
    row arrived needing a manual look.
    """
    labels = ["Publishable matter"] * 6 + ["Non-publishable matter"] * 11
    metric = submission(tmp_path, labels)
    assert metric.value == 17


def test_unnumbered_labels_raise_no_numbering_gap(tmp_path):
    """There are no numbers to be consecutive, so nothing should be claimed."""
    metric = submission(tmp_path, ["Publishable matter"] * 3)
    assert metric.numbering_gap is False
    assert not any("numbered up to" in n for n in metric.notes)


def test_the_matter_summary_heading_is_not_a_matter(tmp_path):
    """The section is introduced by "Publishable matter summary", a heading.

    It sits in a table of its own and would be counted as an eighteenth matter
    if the label were matched loosely once the number became optional.
    """
    labels = ["Publishable matter summary"] + ["Publishable matter"] * 3
    metric = submission(tmp_path, labels)
    assert metric.value == 3


def test_numbered_labels_with_trailing_text_still_count(tmp_path):
    metric = submission(tmp_path, ["Publishable matter 1 (redacted)", "Publishable matter 2"])
    assert metric.value == 2


def test_mixed_numbered_and_unnumbered_labels(tmp_path):
    metric = submission(tmp_path, ["Publishable matter 1", "Publishable matter"])
    assert metric.value == 2


# --- a hash between "Publishable" and "Matter" -----------------------------


def test_a_hash_in_the_matter_label_is_read(tmp_path):
    """One firm types "Publishable #Matter 1" in all but one of its boxes.

    The odd box out was spelled without the hash, so it alone was counted and
    the numbering check reported fourteen missing matters.
    """
    labels = [f"Publishable #Matter {n}" for n in range(1, 15)]
    labels.append("Publishable matter 15")
    metric = submission(tmp_path, labels)
    assert metric.value == 15
    assert metric.numbering_gap is False


def test_the_hash_form_is_counted_as_non_publishable_too(tmp_path):
    metric = submission(tmp_path, ["Non-publishable #Matter 1", "Publishable #Matter 1"])
    assert metric.value == 2
    assert metric.numbering_gap is False


# --- labels punctuated the firm's own way -----------------------------------
#
# A scan of 578 real submissions found eight label shapes the parser read as no
# matter at all. One firm wrote "Publishable - Matter 1" and "Publishable matter
# - 4" for its first six boxes, so 20 matters counted as 14; another wrote
# "Non-publishable matter7" with no space before the number.

PUNCTUATED = [
    "Publishable - Matter 1",
    "Publishable – Matter 1",
    "Publishable matter - 1",
    "Publishable matter: 1",
    "Non-publishable matter1",
    "Non-publishable matter - 1",
    "Non-publishable matter – 1",
    "Non-publishable – Matter 1",
    "Non Publishable matter 1",
]


@pytest.mark.parametrize("label", PUNCTUATED)
def test_a_punctuated_label_is_a_matter(tmp_path, label):
    assert submission(tmp_path, [label]).value == 1


@pytest.mark.parametrize("label", [
    "Publishable matter summary",
    "Publishable - Matter description",
    "Non-publishable matter summary",
])
def test_headings_that_merely_start_like_a_label_are_not_matters(tmp_path, label):
    assert submission(tmp_path, [label, "Publishable matter 1"]).value == 1


def test_a_label_without_its_hyphen_is_still_non_publishable(tmp_path):
    """"Non Publishable matter 1" next to "Publishable matter 1" is no repeat."""
    metric = submission(tmp_path, ["Publishable matter 1", "Non Publishable matter 1"])
    assert metric.numbering_gap is False


def test_the_reported_document_shape(tmp_path):
    """Six publishable boxes with dashes, then non-publishable 7 to 20."""
    labels = ["Publishable - Matter 1", "Publishable - Matter 2", "Publishable – Matter 3"]
    labels += [f"Publishable matter - {n}" for n in (4, 5, 6)]
    labels += [f"Non-publishable matter {n}" for n in range(7, 21)]
    metric = submission(tmp_path, labels)
    assert metric.value == 20
    assert metric.numbering_gap is False


# --- numbering that runs across both kinds ----------------------------------


def test_numbering_straight_through_both_kinds_is_not_a_gap(tmp_path):
    labels = [f"Publishable matter {n}" for n in (1, 2, 3)]
    labels += [f"Non-publishable matter {n}" for n in (4, 5, 6)]
    assert submission(tmp_path, labels).numbering_gap is False


def test_interleaved_numbering_is_not_a_gap(tmp_path):
    labels = ["Publishable matter 1", "Publishable matter 2", "Non-publishable matter 3",
              "Publishable matter 4", "Non-publishable matter 5"]
    assert submission(tmp_path, labels).numbering_gap is False


def test_a_gap_in_both_readings_is_still_reported(tmp_path):
    labels = ["Publishable matter 1", "Publishable matter 2", "Non-publishable matter 4"]
    metric = submission(tmp_path, labels)
    assert metric.numbering_gap is True
    assert "no 3" in " ".join(metric.notes)


def test_a_mistyped_number_gives_a_short_note(tmp_path):
    """One firm typed 230 for 30; the note named all 200 numbers between."""
    labels = [f"Publishable matter {n}" for n in range(1, 30)] + ["Non-publishable matter 230"]
    metric = submission(tmp_path, labels)
    assert metric.value == 30
    assert metric.numbering_gap is True
    note = next(n for n in metric.notes if "numbered up to" in n)
    assert "no 30 to 229" in note
    assert len(note) < 400


def test_missing_numbers_are_written_as_ranges():
    from parsing import _ranges

    assert _ranges([5, 7, 8, 9] + list(range(30, 230))) == "5, 7 to 9, 30 to 229"
    assert _ranges([3, 4]) == "3, 4"
    assert _ranges([14]) == "14"


# --- the non-publishable box worded another way -----------------------------
#
# One firm labelled its last red, non-publishable box "Confidential matter 5"
# after "Non-publishable matter 1" to 4, so 26 matters counted as 25. Others
# wrote "Non-published matter 9" and, by a slip, "Son-publishable matter 3".


@pytest.mark.parametrize("label", [
    "Confidential matter 1",
    "Confidential Matter 1",
    "Confidential - Matter 1",
    "Non-published matter 1",
    "Son-publishable matter 1",
])
def test_another_wording_of_a_non_publishable_box_is_a_matter(tmp_path, label):
    assert submission(tmp_path, [label]).value == 1


def test_a_confidential_box_continues_the_non_publishable_numbering(tmp_path):
    """The reported shape: Non-publishable 1 to 4, then Confidential 5."""
    labels = [f"Publishable matter {n}" for n in range(1, 22)]
    labels += [f"Non-publishable matter {n}" for n in range(1, 5)]
    labels.append("Confidential matter 5")
    metric = submission(tmp_path, labels)
    assert metric.value == 26
    assert metric.numbering_gap is False


def test_confidential_matter_summary_is_not_a_matter(tmp_path):
    assert submission(tmp_path, ["Confidential matter summary", "Publishable matter 1"]).value == 1
