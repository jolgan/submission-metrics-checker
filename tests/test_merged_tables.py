"""Two entries sharing one table.

Word merges adjacent tables into one when the empty paragraph between them is
deleted, which happens readily while a firm edits its submission. The second
entry's label then survives only as an ordinary row part-way down the merged
table. A parser that reads the first row alone counts one where there are two.

A real submission was counted as 19 work matters instead of 20 for exactly this
reason: matters 12 and 13 shared a 54-row table, with 'Publishable matter 13'
sitting at row 31.
"""

import pytest
from docx import Document

from conftest import SAMPLES
from parsing import parse_document

from conftest import EXPECTED

MATTERS = {k: v["matters"] for k, v in EXPECTED.items()}


def submission(tmp_path, name="doc.docx"):
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    return document, tmp_path / name


def merged_matters(document, numbers):
    """One table holding several matters, as Word leaves them after a merge."""
    table = document.add_table(rows=3 * len(numbers), cols=2)
    for position, number in enumerate(numbers):
        top = position * 3
        table.rows[top].cells[0].text = f"Publishable matter {number}"
        table.rows[top + 1].cells[0].text = "Name of client"
        table.rows[top + 1].cells[1].text = "Industry sector"
        table.rows[top + 2].cells[0].text = f"Client {number}"
        table.rows[top + 2].cells[1].text = "Life Sciences"
    return table


# --- matters ----------------------------------------------------------------


def test_two_matters_in_one_table_are_both_counted(tmp_path):
    document, path = submission(tmp_path)
    merged_matters(document, [12, 13])
    document.save(path)

    metric = parse_document(path).metrics["matters"]
    assert metric.value == 2
    assert "Publishable matter 13" in " ".join(metric.evidence)
    assert any("share a table" in n for n in metric.notes)


def test_a_long_merge_counts_every_matter(tmp_path):
    document, path = submission(tmp_path)
    merged_matters(document, [1, 2, 3, 4, 5])
    document.save(path)
    assert parse_document(path).metrics["matters"].value == 5


def test_an_empty_matter_inside_a_merged_table_is_skipped(tmp_path):
    document, path = submission(tmp_path)
    table = document.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Publishable matter 1"
    table.rows[1].cells[0].text = "Name of client"
    table.rows[1].cells[1].text = "Client One"
    table.rows[2].cells[0].text = "Publishable matter 2"   # nothing under it
    table.rows[3].cells[0].text = ""
    document.save(path)

    metric = parse_document(path).metrics["matters"]
    assert metric.value == 1
    assert any("empty matter template" in n for n in metric.notes)


def test_unmerged_matters_still_count_once_each(tmp_path):
    document, path = submission(tmp_path)
    for number in (1, 2, 3):
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = f"Publishable matter {number}"
        table.rows[1].cells[0].text = "Name of client"
        table.rows[1].cells[1].text = f"Client {number}"
    document.save(path)

    metric = parse_document(path).metrics["matters"]
    assert metric.value == 3
    assert not any("share a table" in n for n in metric.notes)


# --- nominations ------------------------------------------------------------


def merged_nominations(document, entries):
    table = document.add_table(rows=3 * len(entries), cols=2)
    for position, (label, name) in enumerate(entries):
        top = position * 3
        table.rows[top].cells[0].text = label
        table.rows[top + 1].cells[0].text = "Name"
        table.rows[top + 1].cells[1].text = "Location"
        table.rows[top + 2].cells[0].text = name
        table.rows[top + 2].cells[1].text = "London"
    return table


def test_two_nominations_in_one_table_are_both_counted(tmp_path):
    document, path = submission(tmp_path)
    merged_nominations(document, [
        ("Partner: leading partner 1", "Alice Adams"),
        ("Partner: leading partner 2", "Bruno Baker"),
    ])
    document.save(path)

    metric = parse_document(path).metrics["lead_partners"]
    assert metric.value == 2
    assert "Bruno Baker" in " ".join(metric.evidence)


def test_a_merged_table_can_hold_different_categories(tmp_path):
    document, path = submission(tmp_path)
    merged_nominations(document, [
        ("Partner: leading partner 1", "Alice Adams"),
        ("Partner: next generation partner 1", "Bruno Baker"),
        ("Associate: leading associate 1", "Cara Clarke"),
    ])
    document.save(path)

    metrics = parse_document(path).metrics
    assert metrics["lead_partners"].value == 1
    assert metrics["next_gen"].value == 1
    assert metrics["associates"].value == 1
    assert "Alice Adams" in " ".join(metrics["lead_partners"].evidence)
    assert "Bruno Baker" in " ".join(metrics["next_gen"].evidence)
    assert "Cara Clarke" in " ".join(metrics["associates"].evidence)


def test_a_nominee_name_does_not_leak_from_the_entry_above(tmp_path):
    """The second entry is empty; it must not inherit the first one's name."""
    document, path = submission(tmp_path)
    table = document.add_table(rows=5, cols=2)
    table.rows[0].cells[0].text = "Partner: leading partner 1"
    table.rows[1].cells[0].text = "Name"
    table.rows[1].cells[1].text = "Location"
    table.rows[2].cells[0].text = "Alice Adams"
    table.rows[3].cells[0].text = "Partner: leading partner 2"
    table.rows[4].cells[0].text = ""
    document.save(path)

    metric = parse_document(path).metrics["lead_partners"]
    assert metric.value == 1, "the empty second entry must not count"
    assert "Alice Adams" in " ".join(metric.evidence)
    assert any("empty nomination template" in n for n in metric.notes)


# --- the samples are unaffected --------------------------------------------


@pytest.mark.parametrize("doc_key,expected", sorted(MATTERS.items()))
def test_sample_matter_counts_are_unchanged(doc_key, expected):
    assert parse_document(SAMPLES[doc_key]).metrics["matters"].value == expected


@pytest.mark.parametrize("doc_key", sorted(MATTERS))
def test_sample_nomination_counts_are_unchanged(doc_key):
    counts = EXPECTED[doc_key]
    expected = (counts["lead_partners"], counts["next_gen"], counts["associates"])
    metrics = parse_document(SAMPLES[doc_key]).metrics
    assert (
        metrics["lead_partners"].value,
        metrics["next_gen"].value,
        metrics["associates"].value,
    ) == expected
