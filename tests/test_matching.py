"""Export loading, filename extraction, matching and the .xlsx writer."""

from io import BytesIO

import openpyxl
import pytest

from conftest import DETAILS, EXPECTED, PORTAL, POWERBI, SAMPLES
from export_writer import GREEN, build_workbook
from matching import (
    build_results,
    difference_records,
    firm_key,
    load_portal,
    load_powerbi,
    normalise,
    results_to_frame,
    to_int,
    uri_to_filename,
)
from parsing import parse_document


@pytest.fixture(scope="module")
def powerbi():
    return load_powerbi(POWERBI)


@pytest.fixture(scope="module")
def portal():
    return load_portal(PORTAL)


@pytest.fixture(scope="module")
def results(powerbi, portal):
    documents = [parse_document(p) for p in SAMPLES.values()]
    return build_results(documents, powerbi, portal)


# --- normalisation ----------------------------------------------------------


def test_normalise_lowercases_trims_and_collapses():
    assert normalise("  Real   Estate  ") == "real estate"


def test_normalise_standardises_spacing_around_gt():
    variants = [
        "Real estate>Commercial property",
        "Real estate > Commercial property",
        "Real estate  >  Commercial property",
        "Real estate>  Commercial property",
    ]
    assert len({normalise(v) for v in variants}) == 1


def test_normalise_handles_none():
    assert normalise(None) == ""


def test_to_int_tolerates_text_numbers_and_blanks():
    assert to_int("21") == 21
    assert to_int(21) == 21
    assert to_int("") is None
    assert to_int(None) is None
    assert to_int("n/a") is None


def test_firm_key_reduces_suffixes_without_merging_distinct_firms():
    assert firm_key("Aldgate Fenwick LLP") == firm_key("Aldgate Fenwick")
    assert firm_key("Aldgate Fenwick") != firm_key("Aldgate Holdings")


# --- File URI ---------------------------------------------------------------


def test_uri_to_filename_keeps_the_docx_extension():
    uri = "https://example.invalid/sites/x/Shared Documents/London/firm_submission.docx"
    assert uri_to_filename(uri) == "firm_submission.docx"


def test_uri_to_filename_decodes_percent_escapes_and_drops_query():
    uri = "https://example.invalid/Shared%20Documents/firm_2027%20Edition.docx?web=1"
    assert uri_to_filename(uri) == "firm_2027 Edition.docx"


def test_uri_without_an_extension_gets_a_docx_suffix():
    assert uri_to_filename("https://example.invalid/x/Firm_submission_2027") == (
        "Firm_submission_2027.docx"
    )


def test_uri_with_a_different_extension_is_not_renamed():
    assert uri_to_filename("https://example.invalid/x/notes.pdf") == "notes.pdf"


def test_portal_export_finds_the_sample_documents(portal):
    for path in SAMPLES.values():
        assert portal.lookup(path.name) is not None


def test_portal_export_lookup_is_case_insensitive(portal):
    assert portal.lookup(SAMPLES["alpha"].name.upper()) is not None


def test_portal_export_reads_excel_hyperlink_targets(tmp_path):
    """pandas sees only the display text of a hyperlink; we need the target."""
    path = tmp_path / "portal-with-hyperlinks.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Firm", "Region", "Practice Area Name", "File URI"])
    sheet.append(["Test Firm", "London", "Banking and finance", "DOCX"])
    sheet.cell(row=2, column=4).hyperlink = (
        "https://example.invalid/sites/x/test-firm_submission.docx?web=1"
    )
    workbook.save(path)

    portal = load_portal(path)
    record = portal.lookup("test-firm_submission.docx")
    assert record is not None
    assert record["firm"] == "Test Firm"
    assert portal.hyperlink_count == 1


# --- Power BI export --------------------------------------------------------


def test_powerbi_drops_filter_and_spacer_rows(powerbi):
    assert powerbi.dropped > 0
    assert all(row.get("FIRM_NAME") for row in powerbi.rows)
    assert not any(
        "Applied filters" in str(row.get("MATCH_STATUS") or "") for row in powerbi.rows
    )


def test_powerbi_rejects_a_workbook_without_the_expected_columns(tmp_path):
    path = tmp_path / "wrong.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["A", "B", "C"])
    workbook.save(path)
    with pytest.raises(ValueError, match="expected Power BI columns"):
        load_powerbi(path)


def test_region_is_part_of_the_key(powerbi):
    details = DETAILS["alpha"]
    assert powerbi.lookup(details["firm"], details["region"], details["practice_area"])
    assert not powerbi.lookup(details["firm"], "Nowhere", details["practice_area"])


# --- matching ---------------------------------------------------------------


def test_all_sample_documents_match_exactly_one_export_row(results):
    assert len(results) == len(SAMPLES)
    for row in results:
        assert row.status == "matched", (row.filename, row.issues)
        assert row.export_row is not None


def test_matched_rows_carry_firm_reference_and_filename(results):
    for row in results:
        assert row.firm_ref
        assert row.filename.endswith(".docx")
        assert row.firm and row.region and row.practice_area


def test_reference_columns_come_from_the_portal_export(results):
    for row in results:
        assert row.reference_source == "Staff Portal export"


def test_practice_area_survives_an_unselected_dropdown(powerbi, portal):
    """Reference data belongs to the portal record, not the document."""
    doc = parse_document(SAMPLES["alpha"])
    doc.region = doc.practice_group = doc.practice_area = doc.dropdown_raw = None
    rows = build_results([doc], powerbi, portal)
    assert rows[0].practice_area == DETAILS["alpha"]["practice_area"]
    assert rows[0].region == DETAILS["alpha"]["region"]
    assert rows[0].status == "matched"


def test_unmatched_document_is_reported_not_dropped(powerbi, portal):
    doc = parse_document(SAMPLES["alpha"])
    doc.filename = "not-in-the-portal-export.docx"
    doc.region = "Nowhere"
    rows = build_results([doc], powerbi, portal)
    assert len(rows) == 1
    assert rows[0].status == "unmatched"
    assert any("No Power BI export row" in i for i in rows[0].issues)


def test_unmatched_row_suggests_the_regions_that_do_exist(powerbi, portal):
    doc = parse_document(SAMPLES["alpha"])
    doc.filename = "not-in-the-portal-export.docx"
    doc.region = "Nowhere"
    rows = build_results([doc], powerbi, portal)
    assert any(DETAILS["alpha"]["region"] in i for i in rows[0].issues)


def test_portal_export_is_authoritative_but_disagreement_is_reported(powerbi, portal):
    doc = parse_document(SAMPLES["alpha"])
    doc.region = "Nowhere"
    rows = build_results([doc], powerbi, portal)
    assert rows[0].region == DETAILS["alpha"]["region"]
    assert rows[0].status == "matched"
    assert any("Region in the document" in i for i in rows[0].issues)


def test_document_missing_from_portal_still_produces_a_row(powerbi):
    from matching import PortalExport

    doc = parse_document(SAMPLES["alpha"])
    rows = build_results([doc], powerbi, PortalExport())
    assert rows[0].filename == SAMPLES["alpha"].name
    assert any("not found in the Staff Portal export" in i for i in rows[0].issues)


def test_unparsed_metric_is_never_treated_as_a_difference(powerbi, portal):
    doc = parse_document(SAMPLES["alpha"])
    doc.metrics["matters"].value = None
    rows = build_results([doc], powerbi, portal)
    assert not rows[0].differs("matters")
    assert "matters" not in [r["Metric"] for r in difference_records(rows)]


def test_the_export_undercounts_nominations_so_corrections_appear(results):
    """The synthetic export mirrors the real one: nomination counts too low."""
    differences = difference_records(results)
    assert differences, "the samples should produce at least one correction"
    assert any(d["Metric"] == "Next generation" for d in differences)


# --- presentation -----------------------------------------------------------


def test_frame_has_the_exact_agreed_columns_in_order(results):
    frame = results_to_frame(results)
    assert len(frame) == len(SAMPLES)
    assert list(frame.columns) == [
        "FIRMREF", "FIRM", "REGION", "PRACTICE AREA", "FILE NAME",
        "MATTERS", "CLIENTS", "NEW", "LPs", "NGs", "LAs",
    ]


def test_show_original_toggle_reveals_the_export_value(results):
    frame = results_to_frame(results, show_original=True)
    row = frame[frame["FILE NAME"].str.contains("alpha")].iloc[0]
    assert row["NGs"] == f"{EXPECTED['alpha']['next_gen']} (was 0)"
    assert row["MATTERS"] == str(EXPECTED["alpha"]["matters"])


# --- workbook ---------------------------------------------------------------


def test_workbook_bakes_green_fills_into_the_changed_cells(results):
    workbook = openpyxl.load_workbook(BytesIO(build_workbook(results)))
    sheet = workbook["Recount"]
    header = [c.value for c in sheet[1]]
    col = header.index("NGs") + 1
    name_col = header.index("FILE NAME") + 1

    filled = 0
    for row in range(2, sheet.max_row + 1):
        if "alpha" in str(sheet.cell(row=row, column=name_col).value):
            cell = sheet.cell(row=row, column=col)
            assert cell.value == EXPECTED["alpha"]["next_gen"]
            assert cell.fill.start_color.rgb == GREEN.start_color.rgb
            filled += 1
    assert filled == 1


def test_workbook_leaves_agreeing_cells_unfilled(results):
    workbook = openpyxl.load_workbook(BytesIO(build_workbook(results)))
    sheet = workbook["Recount"]
    header = [c.value for c in sheet[1]]
    col = header.index("MATTERS") + 1
    name_col = header.index("FILE NAME") + 1
    for row in range(2, sheet.max_row + 1):
        if "alpha" in str(sheet.cell(row=row, column=name_col).value):
            assert sheet.cell(row=row, column=col).fill.start_color.rgb != (
                GREEN.start_color.rgb
            )


def test_workbook_has_changes_and_issues_sheets(results):
    workbook = openpyxl.load_workbook(BytesIO(build_workbook(results)))
    assert workbook.sheetnames == ["Recount", "Changes", "Issues"]
    assert workbook["Changes"].max_row == len(difference_records(results)) + 1


# --- practice areas spelled with different numbers of levels ---------------


def test_a_repeated_trailing_level_is_reduced():
    from matching import practice_area_variants

    forms = practice_area_variants("Finance > Banking and finance > Banking and finance")
    assert "finance > banking and finance" in forms


def test_a_three_level_name_offers_its_parts():
    from matching import practice_area_variants

    forms = practice_area_variants("Transport > Travel > Travel: personal injury")
    assert "transport > travel" in forms
    assert "travel: personal injury" in forms


def test_a_two_level_name_offers_the_bare_area():
    from matching import practice_area_variants

    assert practice_area_variants("Transport > Travel") == {"transport > travel", "travel"}


def test_a_trailing_level_repeated_with_different_punctuation_is_reduced():
    """The portal's table name is not always punctuated like its area name.

    Eleven Brazilian firms file under a practice area named "Electricity (and
    renewable energy)" whose table drops the brackets. Read literally that is a
    third level, and the bare area the Staff Portal page prints is then indexed
    under nothing, so every one of those submissions matches no portal record.
    """
    from matching import practice_area_variants

    forms = practice_area_variants(
        "Energy and natural resources > Electricity (and renewable energy) "
        "> Electricity and renewable energy"
    )
    assert "electricity (and renewable energy)" in forms
    assert "energy and natural resources > electricity (and renewable energy)" in forms


def test_both_spellings_of_a_collapsed_level_are_kept():
    """Either spelling may be the one pasted, so both have to resolve."""
    from matching import practice_area_variants

    forms = practice_area_variants(
        "Energy and natural resources > Electricity (and renewable energy) "
        "> Electricity and renewable energy"
    )
    assert "electricity and renewable energy" in forms


def test_levels_that_differ_by_more_than_punctuation_are_not_collapsed():
    """The guard is punctuation, not a general resemblance between levels.

    "Travel" and "Travel: personal injury" are two levels, not one spelled two
    ways, so the last level has to survive as part of the name rather than
    being dropped as a repeat.
    """
    from matching import practice_area_variants

    forms = practice_area_variants("Transport > Travel > Travel: personal injury")
    assert "transport > travel > travel: personal injury" in forms
    assert "travel > travel: personal injury" in forms
    assert "travel: personal injury" in forms


def test_the_bare_practice_area_of_a_three_level_name_is_offered():
    """The Staff Portal screen prints the area alone, so it must resolve.

    Where the portal's table name is neither a repeat of the area nor derived
    from it - "City focus - Brasilia - Government relations" for the area
    "Government relations" - the area is the middle of three levels and no
    other form names it.
    """
    from matching import practice_area_variants

    forms = practice_area_variants(
        "City focus - Brasilia > Government relations "
        "> City focus - Brasilia - Government relations"
    )
    assert "government relations" in forms


def test_the_practice_group_alone_is_not_offered():
    """A group covers many areas, so on its own it identifies no submission."""
    from matching import practice_area_variants

    forms = practice_area_variants("Transport > Travel > Travel: personal injury")
    assert "transport" not in forms
    assert practice_area_variants("Transport > Travel") == {"transport > travel", "travel"}


def test_an_export_row_written_with_three_levels_still_matches(powerbi, tmp_path):
    """The Power BI export repeats the last level on some rows."""
    import openpyxl

    from matching import POWERBI_REQUIRED

    path = tmp_path / "three-level.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    header = POWERBI_REQUIRED + ["NUM_MATTERS"]
    sheet.append(header)
    sheet.append(["A Firm LLP", "1", "London",
                  "Finance > Banking and finance > Banking and finance", 7])
    workbook.save(path)

    export = load_powerbi(path)
    assert export.lookup("A Firm LLP", "London", "Finance > Banking and finance")
    assert export.lookup("A Firm LLP", "London", "Banking and finance")


# --- the two systems spell a practice area differently ---------------------


def test_the_documents_spelling_is_tried_when_the_portals_finds_nothing(tmp_path):
    """The portal's "Dispute resolution" is the dashboard's fuller name.

    Reference data comes from the portal, but where that spelling matches no
    export row the document's own is worth a try before giving up.
    """
    import openpyxl

    from matching import PortalExport, build_results
    from parsing import parse_document

    path = tmp_path / "powerbi.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["FIRM_NAME", "Firm Ref", "Country", "Practice Area",
                  "# Matters", "# Active Clients", "# New Clients",
                  "# L.Partners", "# Next Gens", "# Associates"])
    sheet.append(["Aldgate Fenwick LLP", "90001", "South East",
                  "Real estate > Commercial property: Thames Valley",
                  1, 1, 1, 1, 1, 1])
    workbook.save(path)
    export = load_powerbi(path)

    doc = parse_document(SAMPLES["alpha"])
    rows = build_results([doc], export, PortalExport())
    assert rows[0].status == "matched"
    assert rows[0].export_row is not None
