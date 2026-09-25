"""Tables that a plain walk of the document body cannot see.

A submission was counted as 19 work matters instead of 20: matter 13 was
skipped while 12 and 14 were counted, and nothing on screen said a table had
been passed over. A table is invisible to a top-level walk when it is wrapped
in a block-level content control, or nested inside another table's cell.
"""

import copy
import pytest
from docx import Document
from docx.oxml.ns import qn

from conftest import SAMPLES
from parsing import parse_document

from conftest import EXPECTED

EXPECTED_MATTERS = {k: v["matters"] for k, v in EXPECTED.items()}


def submission(path):
    """A minimal submission the parser will accept."""
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    return document


def add_matter(document, number):
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = f"Publishable matter {number}"
    table.rows[1].cells[0].text = "Name of client"
    table.rows[1].cells[1].text = f"Client {number}"
    return table


def wrap_in_content_control(table):
    """Put a table inside a block-level w:sdt, as Word does for templates."""
    element = table._tbl
    sdt = element.makeelement(qn("w:sdt"), {})
    content = element.makeelement(qn("w:sdtContent"), {})
    element.addprevious(sdt)
    sdt.append(content)
    content.append(element)


def nest_in_cell(document, table):
    """Move a table inside another table's cell."""
    outer = document.add_table(rows=1, cols=1)
    outer.rows[0].cells[0]._tc.append(copy.deepcopy(table._tbl))
    table._tbl.getparent().remove(table._tbl)


# --- the reported failure ---------------------------------------------------


def test_a_matter_in_a_content_control_is_counted(tmp_path):
    document = submission(tmp_path)
    tables = [add_matter(document, n) for n in (1, 2, 3)]
    wrap_in_content_control(tables[1])
    document.save(tmp_path / "sdt.docx")

    metric = parse_document(tmp_path / "sdt.docx").metrics["matters"]
    assert metric.value == 3, "a wrapped table must not disappear"
    assert "Publishable matter 2" in " ".join(metric.evidence)


def test_a_matter_nested_in_a_cell_is_counted(tmp_path):
    document = submission(tmp_path)
    inner = add_matter(document, 2)
    nest_in_cell(document, inner)
    add_matter(document, 1)
    add_matter(document, 3)
    document.save(tmp_path / "nested.docx")

    metric = parse_document(tmp_path / "nested.docx").metrics["matters"]
    assert metric.value == 3
    assert "Publishable matter 2" in " ".join(metric.evidence)


def test_a_hidden_nomination_is_counted(tmp_path):
    document = submission(tmp_path)
    for number in (1, 2):
        table = document.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = f"Partner: leading partner {number}"
        table.rows[1].cells[0].text = "Name"
        table.rows[1].cells[1].text = "Location"
        table.rows[2].cells[0].text = f"Partner {number}"
        if number == 2:
            wrap_in_content_control(table)
    add_matter(document, 1)
    document.save(tmp_path / "noms.docx")

    metric = parse_document(tmp_path / "noms.docx").metrics["lead_partners"]
    assert metric.value == 2
    assert "Partner 2" in " ".join(metric.evidence)


def test_a_hidden_client_table_is_counted(tmp_path):
    document = submission(tmp_path)
    extra = document.add_table(rows=2, cols=2)
    extra.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    extra.rows[0].cells[1].text = "New client (yes/no)"
    extra.rows[1].cells[0].text = "Another Client Ltd"
    extra.rows[1].cells[1].text = "No"
    wrap_in_content_control(extra)
    add_matter(document, 1)
    document.save(tmp_path / "clients.docx")

    assert parse_document(tmp_path / "clients.docx").metrics["active_clients"].value == 2


# --- and nothing gets counted twice ----------------------------------------


def test_a_table_is_never_yielded_twice(tmp_path):
    """Descending into cells must not re-yield a table already seen."""
    from docx import Document as Open
    from parsing import iter_blocks

    document = submission(tmp_path)
    for n in (1, 2, 3):
        add_matter(document, n)
    document.save(tmp_path / "plain.docx")

    tables = [item for kind, item in iter_blocks(Open(tmp_path / "plain.docx"))
              if kind == "tbl"]
    ids = [id(t._tbl) for t in tables]
    assert len(ids) == len(set(ids))


def test_the_walk_gives_the_same_answer_lazily_as_into_a_list(tmp_path):
    """lxml reuses the id() of a collected proxy for a different element.

    An earlier version skipped tables it thought it had already seen, which only
    went wrong when the walk was consumed lazily - holding the results in a list
    kept the proxies alive and hid it. The walk must not depend on how it is
    consumed.
    """
    from docx import Document as Open
    from parsing import iter_blocks

    document = submission(tmp_path)
    for n in range(1, 16):
        add_matter(document, n)
    path = tmp_path / "many.docx"
    document.save(path)

    eager = [label for kind, item in iter_blocks(Open(path)) if kind == "tbl"
             for label in [item.rows[0].cells[0].text]]
    lazy = []
    for kind, item in ((k, i) for k, i in iter_blocks(Open(path)) if k == "tbl"):
        lazy.append(item.rows[0].cells[0].text)

    assert eager == lazy
    assert len(eager) == 17, "firm name, client table and fifteen matters"


@pytest.mark.parametrize("doc_key,expected", sorted(EXPECTED_MATTERS.items()))
def test_the_samples_are_unchanged(doc_key, expected):
    """The deeper walk must not find anything new in the known-good samples."""
    doc = parse_document(SAMPLES[doc_key])
    assert doc.metrics["matters"].value == expected


@pytest.mark.parametrize("doc_key", sorted(EXPECTED_MATTERS))
def test_sample_client_counts_are_unchanged(doc_key):
    doc = parse_document(SAMPLES[doc_key])
    assert doc.metrics["active_clients"].value == EXPECTED[doc_key]["active_clients"]
