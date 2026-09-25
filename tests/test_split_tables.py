"""Labels split away from the body they introduce.

The counterpart of the merged-table problem. Word also splits one table into
two, leaving the label alone in a table of its own while its body follows as a
table carrying no label. One firm's submissions had nine matters in that shape:
read literally each label was an empty template and each body belonged to
nothing, so 22 matters counted as 13 and the numbering check then reported the
missing ones as a gap the firm had left.

A template that is genuinely empty is a different thing and must still be
skipped, so the two are told apart by where the label sits: a split leaves the
label as the last row of its table, while an empty template keeps its skeleton
rows underneath it.
"""

import pytest
from docx import Document

from parsing import parse_document

FIELDS = ("Name of client", "Industry sector")


def start(document):
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"


def whole_matter(document, label, client="A Client Ltd"):
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = label
    table.rows[1].cells[0].text = FIELDS[0]
    table.rows[1].cells[1].text = FIELDS[1]
    table.rows[2].cells[0].text = client


def split_matter(document, label, client="A Client Ltd"):
    """Label alone in its own table, body following with no label."""
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = label
    body = document.add_table(rows=2, cols=2)
    body.rows[0].cells[0].text = FIELDS[0]
    body.rows[0].cells[1].text = FIELDS[1]
    body.rows[1].cells[0].text = client


def empty_matter(document, label):
    """A template left in place: rows under the label, nothing filled in."""
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = label


def matters(tmp_path, build, name="doc.docx"):
    document = Document()
    start(document)
    build(document)
    path = tmp_path / name
    document.save(path)
    return parse_document(path).metrics["matters"]


def test_a_split_matter_is_counted(tmp_path):
    def build(document):
        whole_matter(document, "Publishable matter 1")
        split_matter(document, "Publishable matter 2", "B Client Ltd")
    metric = matters(tmp_path, build)
    assert metric.value == 2


def test_the_split_is_reported(tmp_path):
    def build(document):
        split_matter(document, "Publishable matter 1")
    metric = matters(tmp_path, build)
    assert any("label in a table of their own" in n for n in metric.notes)


def test_a_split_matter_closes_the_numbering_gap(tmp_path):
    """The reported shape: 15 numbered, 9 of them split, one gap claimed."""
    def build(document):
        for n in range(1, 16):
            if n in (2, 5, 6, 7, 8, 9, 10, 11, 12):
                split_matter(document, f"Publishable #Matter {n}")
            else:
                whole_matter(document, f"Publishable #Matter {n}")
    metric = matters(tmp_path, build)
    assert metric.value == 15
    assert metric.numbering_gap is False


def test_an_empty_template_is_still_skipped(tmp_path):
    """Skeleton rows under the label mean an empty box, not a split one."""
    def build(document):
        whole_matter(document, "Publishable matter 1")
        empty_matter(document, "Publishable matter 2")
    metric = matters(tmp_path, build)
    assert metric.value == 1
    assert any("empty matter template" in n for n in metric.notes)


def test_a_label_followed_by_another_label_adopts_nothing(tmp_path):
    """Two bare labels in a row: the second is not the first one's body."""
    def build(document):
        document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Publishable matter 1"
        whole_matter(document, "Publishable matter 2")
    metric = matters(tmp_path, build)
    assert metric.value == 1


def test_a_body_is_not_adopted_twice(tmp_path):
    def build(document):
        document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Publishable matter 1"
        split_matter(document, "Publishable matter 2")
    metric = matters(tmp_path, build)
    assert metric.value == 1


def test_a_split_nomination_is_counted(tmp_path):
    document = Document()
    start(document)
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Partner: leading partner 1"
    body = document.add_table(rows=2, cols=2)
    body.rows[0].cells[0].text = "Name"
    body.rows[0].cells[1].text = "Location"
    body.rows[1].cells[0].text = "A Partner"
    path = tmp_path / "nom.docx"
    document.save(path)
    metric = parse_document(path).metrics["lead_partners"]
    assert metric.value == 1
    assert metric.evidence == ["Partner: leading partner 1 - A Partner"]


def test_an_empty_nomination_template_is_still_skipped(tmp_path):
    document = Document()
    start(document)
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Partner: leading partner 1"
    table.rows[1].cells[0].text = "Name"
    table.rows[1].cells[1].text = "Location"
    path = tmp_path / "nom.docx"
    document.save(path)
    metric = parse_document(path).metrics["lead_partners"]
    assert metric.value == 0
    assert any("empty nomination template" in n for n in metric.notes)
