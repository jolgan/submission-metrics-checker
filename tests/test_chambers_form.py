"""Submissions written on Chambers' form instead of ours.

A few firms answer on a competitor's submission form. It asks for the same
things in a different shape: one row per lawyer with a "Partner Y/N" column in
place of a box per nomination, and client lists headed "D0 - PUBLISHABLE
CLIENTS" and "E0 - CONFIDENTIAL CLIENTS" with a numbering column in front of
the name.

Read as though it were our form, the nominations came out as a silent zero and
the clients as uncountable. Legal 500 reads Y as a leading partner and N as a
leading associate, so the figures can be had - but nothing about this document
can be assumed, so every row is flagged whatever the count.
"""

import pytest
from docx import Document

from parsing import parse_document

LAWYERS = [
    ("Alejandro Sainz", "Y"),
    ("Daniel Diaz", "Y"),
    ("Santiago Alessio-Robles", "Y"),
    ("Estefania Sierra", "N"),
    ("Luis Estrella", "N"),
]


def chambers_form(tmp_path, lawyers=LAWYERS, publishable=(("A Client Ltd", "Yes"),),
                  confidential=(("B Client Ltd", "No"),), marker=True):
    document = Document()
    document.add_paragraph("SUBMISSION FORM")
    if marker:
        document.add_paragraph(
            "Please upload submissions online at: https://myaccount.chambers.com")

    lawyer_table = document.add_table(rows=2 + len(lawyers), cols=4)
    lawyer_table.rows[0].cells[0].text = (
        "B9 Information regarding Ranked and Unranked lawyers (including "
        "associates) in this practice area.")
    for index, heading in enumerate(
            ["Name", "Comments or Web Link", "Partner Y/N", "Ranked Y/N"]):
        lawyer_table.rows[1].cells[index].text = heading
    for index, (name, partner) in enumerate(lawyers, start=2):
        lawyer_table.rows[index].cells[0].text = name
        lawyer_table.rows[index].cells[2].text = partner

    for label, rows in (
        ("D0 - PUBLISHABLE CLIENTS - List of this department's clients.", publishable),
        ("E0 - CONFIDENTIAL CLIENTS - List of this department's clients.", confidential),
    ):
        table = document.add_table(rows=2 + len(rows), cols=3)
        table.rows[0].cells[0].text = label
        table.rows[1].cells[1].text = "Name of Client"
        table.rows[1].cells[2].text = "New Client (Yes/No)"
        for index, (name, answer) in enumerate(rows, start=2):
            table.rows[index].cells[0].text = str(index - 1)
            table.rows[index].cells[1].text = name
            table.rows[index].cells[2].text = answer

    matter = document.add_table(rows=2, cols=2)
    matter.rows[0].cells[0].text = "Publishable Matter 1"
    matter.rows[1].cells[0].text = "A Client Ltd"

    path = tmp_path / "chambers.docx"
    document.save(path)
    return parse_document(path)


def test_the_form_is_recognised(tmp_path):
    assert chambers_form(tmp_path).chambers is True


def test_a_partner_is_a_leading_partner(tmp_path):
    metrics = chambers_form(tmp_path).metrics
    assert metrics["lead_partners"].value == 3


def test_a_non_partner_is_a_leading_associate(tmp_path):
    metrics = chambers_form(tmp_path).metrics
    assert metrics["associates"].value == 2


def test_both_client_lists_are_counted(tmp_path):
    """Publishable and confidential, as on our own form."""
    doc = chambers_form(
        tmp_path,
        publishable=(("A Client Ltd", "Yes"), ("B Client Ltd", "No")),
        confidential=(("C Client Ltd", "Yes"),),
    )
    assert doc.metrics["active_clients"].value == 3
    assert doc.metrics["new_clients"].value == 2


def test_unnamed_client_rows_are_skipped(tmp_path):
    """The form ships numbered rows whether or not they are filled in."""
    doc = chambers_form(
        tmp_path,
        publishable=(("A Client Ltd", "Yes"), ("", ""), ("", "")),
        confidential=(("C Client Ltd", "No"),),
    )
    assert doc.metrics["active_clients"].value == 2


def test_every_figure_is_flagged(tmp_path):
    """Whatever the count, this document always wants checking by eye."""
    metrics = chambers_form(tmp_path).metrics
    for key in ("matters", "active_clients", "new_clients",
                "lead_partners", "associates"):
        assert metrics[key].needs_check is True, key


def test_the_row_says_which_template(tmp_path):
    from matching import ResultRow, review_note

    row = ResultRow(document=chambers_form(tmp_path), filename="chambers.docx")
    assert review_note(row, {}) == "Chambers template used - check every figure"


def test_next_generation_is_not_claimed_as_zero(tmp_path):
    """The form has no next generation section, so there is nothing to count."""
    metric = chambers_form(tmp_path).metrics["next_gen"]
    assert metric.value is None


def test_a_lawyer_answering_neither_is_reported(tmp_path):
    doc = chambers_form(tmp_path, lawyers=[("Alejandro Sainz", "Y"),
                                           ("Ana Nunez", "maybe")])
    assert doc.metrics["lead_partners"].value == 1
    assert doc.metrics["associates"].value == 0
    assert "neither Y nor N" in " ".join(doc.metrics["associates"].notes)


def test_our_own_form_is_not_mistaken_for_theirs(tmp_path):
    """Recognised by Chambers' own section codes, not by what is missing."""
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    path = tmp_path / "ours.docx"
    document.save(path)
    assert parse_document(path).chambers is False
