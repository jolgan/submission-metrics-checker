"""Rows in a client table that are not clients.

Firms write things into the client list that are not client names: a heading
dividing the list into sections ("ISSUERS", "Initial Purchasers"), or a note
standing in for the list altogether ("All clients are confidential", "ALL
CLIENTS ARE PUBLISHABLE"). Counted literally these inflate the client figure,
and nothing on screen said so - one was only found because the same document
happened to be flagged for its matter numbering.

The text cannot be matched directly: client names vary by country and
language, and a rule written against English phrases would miss the next one
and might drop a real client. What these rows have in common is structural -
they answer the new-client column with nothing, while the rows around them
answer it. That is also what a client whose answer was simply missed looks
like, so the row is counted and flagged rather than dropped.
"""

import pytest
from docx import Document

from parsing import parse_document


def submission(tmp_path, publishable, non_publishable=(), name="doc.docx"):
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"

    for label, rows in (
        ("Clients: publishable clients", publishable),
        ("Clients: non-publishable clients", non_publishable),
    ):
        if not rows:
            continue
        document.add_paragraph(label)
        table = document.add_table(rows=len(rows) + 1, cols=2)
        table.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
        table.rows[0].cells[1].text = "New client (yes/no)"
        for index, (client, answer) in enumerate(rows, start=1):
            table.rows[index].cells[0].text = client
            table.rows[index].cells[1].text = answer
    path = tmp_path / name
    document.save(path)
    return parse_document(path).metrics


def test_a_section_heading_inside_the_list_is_flagged(tmp_path):
    """The reported shape: 63 clients counted as 66."""
    rows = [
        ("ISSUERS", ""),
        ("Acosta Verde, S.A.B. de C.V.", "No"),
        ("Alsea, S.A.B. de C.V.", "No"),
        ("Initial Purchasers", ""),
        ("Banco BTG Pactual S.A.", "No"),
    ]
    metrics = submission(tmp_path, rows)
    assert metrics["active_clients"].needs_check is True
    note = " ".join(metrics["active_clients"].notes)
    assert "ISSUERS" in note and "Initial Purchasers" in note


def test_a_note_standing_in_for_the_whole_list_is_flagged(tmp_path):
    """"All clients are confidential" is the only row in its table."""
    metrics = submission(
        tmp_path,
        [("All clients are confidential", "")],
        [("Merck Sharp & Dohme Comercializadora", "No"), ("Pfizer", "Yes")],
    )
    assert metrics["active_clients"].needs_check is True
    assert "All clients are confidential" in " ".join(metrics["active_clients"].notes)


def test_a_whole_column_left_blank_is_not_flagged(tmp_path):
    """Two firms answer nothing at all. That is a habit, not a heading.

    Flagging every blank answer would flag those firms entirely, which is the
    noise that makes a warning worth ignoring.
    """
    rows = [("Codelco", ""), ("CVC Capital Partners", ""), ("Promigas", "")]
    metrics = submission(tmp_path, rows)
    assert metrics["active_clients"].value == 3
    assert metrics["active_clients"].needs_check is False


def test_a_fully_answered_table_is_not_flagged(tmp_path):
    rows = [("Harbour Wells Group", "No"), ("Trenton Mills plc", "Yes")]
    metrics = submission(tmp_path, rows)
    assert metrics["active_clients"].needs_check is False


def test_the_flagged_rows_are_still_counted(tmp_path):
    """Dropping them would risk losing a client whose answer was missed."""
    rows = [("ISSUERS", ""), ("Acosta Verde, S.A.B. de C.V.", "No")]
    metrics = submission(tmp_path, rows)
    assert metrics["active_clients"].value == 2


def test_a_client_whose_answer_was_missed_is_flagged_too(tmp_path):
    """The same signal, read the other way round, and worth saying."""
    rows = [("Kuraray", ""), ("Harbour Wells Group", "No")]
    metrics = submission(tmp_path, rows)
    assert metrics["active_clients"].needs_check is True
    assert "Kuraray" in " ".join(metrics["active_clients"].notes)


def test_the_review_note_names_clients_not_nominations(tmp_path):
    """The row-level note had only nomination wording for this flag."""
    from matching import review_note, ResultRow

    rows = [("ISSUERS", ""), ("Acosta Verde, S.A.B. de C.V.", "No")]
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    table.rows[0].cells[1].text = "New client (yes/no)"
    for index, (client, answer) in enumerate(rows, start=1):
        table.rows[index].cells[0].text = client
        table.rows[index].cells[1].text = answer
    path = tmp_path / "note.docx"
    document.save(path)

    row = ResultRow(document=parse_document(path), filename="note.docx")
    assert "may not be a client" in review_note(row, {})
