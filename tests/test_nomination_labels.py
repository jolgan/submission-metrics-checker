"""Nomination labels a firm has worded its own way.

One submission used both 'Associate: leading associate 1' and
'Associate : leading counsel 1' - the second with a space before the colon and
a description the parser did not know - so associates counted 1 instead of 3.

The role before the colon is the reliable part. An 'Associate:' label is an
associate whatever the firm calls the role; a 'Partner:' label still has to be
sorted into leading or next generation, and where the description says neither,
it is reported rather than guessed at.
"""

import pytest
from docx import Document

from conftest import SAMPLES
from parsing import UNCLASSIFIED, classify_nomination, parse_document

ASSOCIATE_LABELS = [
    "Associate: leading associate 1",
    "Associate: rising star 2",
    "Associate: leading counsel 1",
    "Associate : leading counsel 1",
    "associate:  senior associate 4",
    "ASSOCIATE: Leading Associate 5",
]

LEAD_LABELS = [
    "Partner: leading partner 1",
    "Partner: leading individual 1",
    "Partner : leading partner 2",
    "Partner: leading counsel 1",
    "Partner: senior partner 3",
]

NEXT_GEN_LABELS = [
    "Partner: next generation 1",
    "Partner: next generation partner 2",
    "Partner: to become next generation partner 1",
    "Partner: nextgen partner 4",
]


@pytest.mark.parametrize("label", ASSOCIATE_LABELS)
def test_any_associate_label_is_an_associate(label):
    assert classify_nomination(label) == "associates"


@pytest.mark.parametrize("label", LEAD_LABELS)
def test_leading_partner_descriptions(label):
    assert classify_nomination(label) == "lead_partners"


@pytest.mark.parametrize("label", NEXT_GEN_LABELS)
def test_next_generation_descriptions(label):
    assert classify_nomination(label) == "next_gen"


def test_next_generation_wins_over_leading():
    """A label naming both must not be filed as a leading partner."""
    assert classify_nomination("Partner: leading next generation partner 1") == "next_gen"


@pytest.mark.parametrize(
    "label", ["Name", "Publishable matter 1", "Partner", "Associate", "Partners: two"]
)
def test_non_nomination_labels_are_not_claimed(label):
    assert classify_nomination(label) is None


def test_an_unknown_partner_description_is_not_guessed():
    assert classify_nomination("Partner: something else entirely 1") == UNCLASSIFIED


# --- end to end -------------------------------------------------------------


def submission(tmp_path, labels, name="doc.docx"):
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"

    for index, label in enumerate(labels, start=1):
        table = document.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = label
        table.rows[1].cells[0].text = "Name"
        table.rows[1].cells[1].text = "Location"
        table.rows[2].cells[0].text = f"Nominee {index}"
        table.rows[2].cells[1].text = "London"
    path = tmp_path / name
    document.save(path)
    return parse_document(path).metrics


def test_the_reported_mix_counts_three_associates(tmp_path):
    metrics = submission(tmp_path, [
        "Associate: leading associate 1",
        "Associate : leading counsel 1",
        "Associate: leading counsel 2",
    ])
    assert metrics["associates"].value == 3
    assert metrics["lead_partners"].value == 0


def test_an_unclassified_partner_is_reported_and_left_uncounted(tmp_path):
    metrics = submission(tmp_path, [
        "Partner: leading partner 1",
        "Partner: something else entirely 1",
    ])
    assert metrics["lead_partners"].value == 1, "the odd one is not assumed to be leading"
    assert metrics["next_gen"].value == 0
    assert metrics["lead_partners"].needs_check is True
    assert metrics["next_gen"].needs_check is True
    note = next(n for n in metrics["lead_partners"].notes if "neither leading" in n)
    assert "something else entirely 1" in note


def test_an_empty_unclassified_label_is_not_reported(tmp_path):
    """Nothing filled in, so nothing to decide."""
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Partner: something else entirely 1"
    table.rows[1].cells[0].text = "Name"
    table.rows[2].cells[0].text = ""
    path = tmp_path / "empty.docx"
    document.save(path)

    metrics = parse_document(path).metrics
    assert metrics["lead_partners"].needs_check is False


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_the_samples_are_unchanged(doc_key):
    from conftest import EXPECTED

    counts = EXPECTED[doc_key]
    expected = (counts["lead_partners"], counts["next_gen"], counts["associates"])
    metrics = parse_document(SAMPLES[doc_key]).metrics
    assert (
        metrics["lead_partners"].value,
        metrics["next_gen"].value,
        metrics["associates"].value,
    ) == expected
    assert not any(
        metrics[k].needs_check for k in ("lead_partners", "next_gen", "associates")
    )


# --- labels with no role prefix, and dashes in place of colons -------------
#
# One submission labelled its boxes 'Next Generation Partner 2' and
# 'Associate - Rising Star 1'. The first has no role in front of it at all; the
# second separates role from description with a dash. Next generation counted
# one where there were two, and all three associates were missed - which at
# least raised "a heading for this category is present but no matching
# nomination table was found" rather than reporting a confident zero.

DASH_SEPARATED = [
    ("Associate – Rising Star 1", "associates"),
    ("Associate — Leading Associate 2", "associates"),
    ("Associate - rising star 3", "associates"),
    ("Partner – leading partner 1", "lead_partners"),
    ("Partner - next generation 2", "next_gen"),
]

BARE_CATEGORY = [
    ("Next Generation Partner 2", "next_gen"),
    ("Next generation 4", "next_gen"),
    ("Next Generation Partners", "next_gen"),
    ("Rising Star 1", "associates"),
    ("Rising Stars 2", "associates"),
    ("Leading Associate 3", "associates"),
    ("Leading Partner 2", "lead_partners"),
    ("Senior Individual 1", "lead_partners"),
]

# Phrases that appear inside matter tables and must not be claimed.
NOT_LABELS = [
    "Lead partner(s)",
    "Alex Morgan (Partner)",
    "Sam Turner, Senior Associate",
    "Start date",
    "Name",
    "Other key team members",
    "Partnership",
]


@pytest.mark.parametrize("label,category", DASH_SEPARATED)
def test_a_dash_separates_role_from_description(label, category):
    assert classify_nomination(label) == category


@pytest.mark.parametrize("label,category", BARE_CATEGORY)
def test_a_label_naming_only_the_category(label, category):
    assert classify_nomination(label) == category


@pytest.mark.parametrize("label", NOT_LABELS)
def test_prose_and_matter_headings_are_not_nominations(label):
    assert classify_nomination(label) is None


def test_the_reported_document_shape(tmp_path):
    """Two next generation boxes, one prefixed and one not."""
    metrics = submission(tmp_path, [
        "Partner: next generation 1",
        "Next Generation Partner 2",
        "Associate – Rising Star 1",
        "Associate – Rising Star 2",
    ])
    assert metrics["next_gen"].value == 2
    assert metrics["associates"].value == 2
    assert metrics["lead_partners"].value == 0


# --- counsel, which the form puts in the leading associates section --------


@pytest.mark.parametrize("label", [
    "Counsel: leading counsel 1",
    "Counsel : leading counsel 2",
    "counsel - leading counsel 3",
    "COUNSEL: Leading Counsel 1",
])
def test_a_counsel_label_is_a_leading_associate(label):
    """The form says so itself: "Note that this section can include counsel."

    Counsel is a role of its own, so the box is labelled "Counsel:" rather
    than "Associate:" and was classified as no category at all, dropping the
    nomination from the count in silence.
    """
    assert classify_nomination(label) == "associates"


def test_a_bare_leading_counsel_label_is_an_associate():
    assert classify_nomination("leading counsel 2") == "associates"


def test_counsel_in_a_partner_description_is_not_reassigned():
    """The role before the colon decides, not the word appearing anywhere."""
    assert classify_nomination("Partner: counsel to the board") == UNCLASSIFIED


def test_the_reported_counsel_document_shape(tmp_path):
    """Three associate boxes, one of them counsel, two numbered alike.

    The firm numbered both associate boxes 2. They name different people, so
    both are nominations and the count is three, not two.
    """
    metrics = submission(tmp_path, [
        "Counsel: leading counsel 1",
        "Associate: leading associate 2",
        "Associate: leading associate 2",
    ])
    assert metrics["associates"].value == 3
    assert metrics["lead_partners"].value == 0
    assert metrics["next_gen"].value == 0


# --- the box number placed before the colon --------------------------------


@pytest.mark.parametrize("label,category", [
    ("Partner 1: leading individual", "lead_partners"),
    ("Partner 2: leading individual", "lead_partners"),
    ("Partner 1: next generation", "next_gen"),
    ("Associate 3: rising star", "associates"),
    ("Counsel 2: leading counsel", "associates"),
])
def test_a_number_between_the_role_and_the_colon(label, category):
    """One firm numbers the box before the colon, not after the description.

    "Partner 1: leading individual" matched nothing, so all three of that
    firm's submissions reported no leading partners at all.
    """
    assert classify_nomination(label) == category


def test_a_plural_role_is_not_a_nomination_label():
    """"Number of PARTNERS in the team" heads a table of its own."""
    assert classify_nomination("Number of PARTNERS in the team") is None


# --- the role written plural ------------------------------------------------


@pytest.mark.parametrize("label,category", [
    ("Associates: leading associate 1", "associates"),
    ("Associates: leading associate 5", "associates"),
    ("Counsels: leading counsel 1", "associates"),
    ("Partners: leading partner 2", "lead_partners"),
    ("Partners: next generation 2", "next_gen"),
])
def test_a_plural_role_is_the_same_role(label, category):
    """One firm carries the section heading's plural onto every box.

    "Associates: leading associate 1" matched no label at all, so five
    nominations were dropped and the category reported as uncountable.
    """
    assert classify_nomination(label) == category


def test_a_plural_associate_is_not_filed_as_a_partner():
    """The guard on the fix above.

    Allowing the plural without also reading it as the same role left
    "Associates:" falling through to the partner branch, where "leading" in
    the description made it a leading partner - turning five dropped
    associates into five invented partners, which is worse than counting none.
    """
    for label in ("Associates: leading associate 1", "Counsels: leading counsel 1"):
        assert classify_nomination(label) != "lead_partners"


def test_a_plural_role_without_a_colon_is_not_a_label():
    for label in ("Number of PARTNERS in the team", "Partners in the team"):
        assert classify_nomination(label) is None


def test_the_reported_plural_document_shape(tmp_path):
    metrics = submission(tmp_path, [
        "Partner: leading partner 1",
        "Associates: leading associate 1",
        "Associates: leading associate 2",
    ])
    assert metrics["associates"].value == 2
    assert metrics["lead_partners"].value == 1


def test_a_plural_role_with_no_nomination_description_is_not_a_label():
    """"Partners: two" is a summary row that happens to share the shape."""
    assert classify_nomination("Partners: two") is None
    assert classify_nomination("Associates: 14") is None


def test_a_singular_role_with_an_invented_description_is_still_reported():
    """The looser reading stays on the singular form, so nothing is dropped."""
    assert classify_nomination("Partner: star performer 1") == UNCLASSIFIED
