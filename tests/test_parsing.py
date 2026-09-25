"""Counting, against the synthetic submissions.

Expected figures come from samples/synthetic/manifest.json, written by the same
script that builds the documents. The three differ on purpose, so between them
they cover the shapes the parser exists to cope with.
"""

import pytest

from conftest import DETAILS, EXPECTED, SAMPLES
from parsing import METRIC_KEYS, clean, parse_document, split_dropdown


@pytest.fixture(scope="module")
def parsed():
    return {key: parse_document(path) for key, path in SAMPLES.items()}


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
@pytest.mark.parametrize("metric", METRIC_KEYS)
def test_metric_counts(parsed, doc_key, metric):
    assert parsed[doc_key].metrics[metric].value == EXPECTED[doc_key][metric]


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_documents_parse_without_error(parsed, doc_key):
    assert parsed[doc_key].ok, parsed[doc_key].error


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_no_metric_is_unparsed(parsed, doc_key):
    assert parsed[doc_key].unparsed_metrics == []


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_header_fields(parsed, doc_key):
    doc = parsed[doc_key]
    details = DETAILS[doc_key]
    assert doc.firm == details["firm"]
    assert doc.region == details["region"]
    assert doc.concatenated_practice_area == details["practice_area"]


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_every_count_carries_evidence(parsed, doc_key):
    """A number the checker cannot trace back to the document is not usable."""
    for key, metric in parsed[doc_key].metrics.items():
        if metric.value:
            assert metric.evidence, f"{doc_key}/{key} counted with no evidence"


# --- the awkward shapes each document carries ------------------------------


def test_duplicate_label_numbers_are_counted_separately(parsed):
    """alpha has two nominations both labelled 'next generation partner 2'."""
    metric = parsed["alpha"].metrics["next_gen"]
    assert metric.value == 2
    assert len(metric.evidence) == 2
    assert "Desmond Fairlie" in " ".join(metric.evidence)
    assert "Priya Ellwood" in " ".join(metric.evidence)


def test_a_blank_answer_counts_as_not_new(parsed):
    metric = parsed["alpha"].metrics["new_clients"]
    assert any("blank answer" in n for n in metric.notes)


def test_merged_matter_tables_are_counted_separately(parsed):
    """beta has matters 2 and 3 sharing one table."""
    metric = parsed["beta"].metrics["matters"]
    assert metric.value == EXPECTED["beta"]["matters"]
    assert any("share a table" in n for n in metric.notes)
    assert "Publishable matter 3" in metric.evidence


def test_a_nomination_inside_a_content_control_is_found(parsed):
    metric = parsed["beta"].metrics["next_gen"]
    assert metric.value == 1
    assert "Tobias Wren" in " ".join(metric.evidence)


def test_a_firms_own_nomination_wording_is_understood(parsed):
    """beta uses 'Associate : leading counsel 2', spacing and all."""
    metric = parsed["beta"].metrics["associates"]
    assert metric.value == 2
    assert "Emeka Braithwaite" in " ".join(metric.evidence)


def test_a_repeated_header_row_inside_a_client_table_is_skipped(parsed):
    notes = " ".join(parsed["gamma"].metrics["new_clients"].notes)
    assert "repeated header row" in notes


def test_new_and_existing_vocabulary_is_substituted_and_flagged(parsed):
    metric = parsed["gamma"].metrics["new_clients"]
    assert metric.value == EXPECTED["gamma"]["new_clients"]
    assert any("'New'/'Existing'" in n for n in metric.notes)


def test_a_qualified_yes_counts(parsed):
    notes = " ".join(parsed["gamma"].metrics["new_clients"].notes)
    assert "qualified answer" in notes


def test_not_applicable_is_not_a_client(parsed):
    notes = " ".join(parsed["gamma"].metrics["active_clients"].notes)
    assert "N/A" in notes


def test_an_empty_nomination_template_is_skipped(parsed):
    metric = parsed["gamma"].metrics["next_gen"]
    assert metric.value == 0
    assert any("empty nomination template" in n for n in metric.notes)


def test_a_numbering_gap_is_reported(parsed):
    """gamma skips publishable matter 3, as firms do when they delete one."""
    metric = parsed["gamma"].metrics["matters"]
    assert metric.numbering_gap is True
    assert any("no 3" in n for n in metric.notes)


def test_the_matter_summary_table_is_not_counted_as_a_matter(parsed):
    for doc in parsed.values():
        for line in doc.metrics["matters"].evidence:
            assert "summary" not in line.lower()


def test_an_absent_category_reports_zero_not_unparsed(parsed):
    metric = parsed["alpha"].metrics["associates"]
    assert metric.value == 0
    assert not metric.is_unparsed


# --- unit-level helpers -----------------------------------------------------


def test_split_dropdown_keeps_hyphens_and_colons_inside_names():
    parts = split_dropdown("South East - Real estate - Commercial property: Berks, Oxon")
    assert parts == ["South East", "Real estate", "Commercial property: Berks, Oxon"]


def test_split_dropdown_does_not_split_on_an_internal_hyphen():
    parts = split_dropdown("London - Employment - Pensions (non-contentious)")
    assert parts[2] == "Pensions (non-contentious)"


def test_clean_collapses_whitespace():
    assert clean("  a \n b\tc ") == "a b c"


def test_an_unreadable_file_is_an_error_not_a_zero(tmp_path):
    bad = tmp_path / "not-a-document.docx"
    bad.write_bytes(b"this is not a docx")
    doc = parse_document(bad)
    assert not doc.ok
    assert all(m.is_unparsed for m in doc.metrics.values())
    assert not any(m.value == 0 for m in doc.metrics.values())


# --- publications outside the UK -------------------------------------------


def test_a_two_part_dropdown_is_region_and_practice_area(tmp_path):
    """Latin America has no practice group, so the dropdown has two parts."""
    from docx import Document
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm"
    prompt = document.add_paragraph("EITHER select Practice Area from this drop-own list")
    run = prompt.add_run()
    sdt = run._r.makeelement(qn("w:sdt"), {})
    props = run._r.makeelement(qn("w:sdtPr"), {})
    alias = run._r.makeelement(qn("w:alias"), {qn("w:val"): "Select Practice "})
    props.append(alias)
    content = run._r.makeelement(qn("w:sdtContent"), {})
    text_run = run._r.makeelement(qn("w:r"), {})
    text = run._r.makeelement(qn("w:t"), {})
    text.text = "Argentina - Dispute resolution: litigation and arbitration"
    text_run.append(text)
    content.append(text_run)
    sdt.append(props)
    sdt.append(content)
    prompt._p.append(sdt)

    document.add_paragraph("Clients: publishable clients")
    clients = document.add_table(rows=2, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[1].cells[0].text = "A Client"
    clients.rows[1].cells[1].text = "Yes"
    path = tmp_path / "latam.docx"
    document.save(path)

    doc = parse_document(path)
    assert doc.region == "Argentina"
    assert doc.practice_group is None
    assert doc.practice_area == "Dispute resolution: litigation and arbitration"
    assert doc.concatenated_practice_area == (
        "Dispute resolution: litigation and arbitration"
    )
    assert not any("did not split" in w for w in doc.warnings)
