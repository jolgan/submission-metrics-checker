"""Packages that are inconsistent about the capitalisation of their own parts.

Seven submissions from one firm could not be opened at all: a part stored as
'customXml/item5.xml' was referenced as 'customXML/item5.xml', and two zip
entries spelled their name one way in the archive index and another in the
entry itself. Word opens these; Python's zipfile refuses, and the documents
were lost rather than counted.
"""

import struct
import zipfile
from io import BytesIO

import pytest
from docx import Document

from conftest import SAMPLES
from docx_repair import repair
from parsing import parse_document


def rebuild(source, rename=None, repoint=None):
    """Copy a .docx, optionally renaming a part or misdirecting a reference."""
    rename = rename or {}
    repoint = repoint or {}
    out = BytesIO()
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            for old, new in repoint.items():
                data = data.replace(old.encode(), new.encode())
            dst.writestr(rename.get(info.filename, info.filename), data)
    return out.getvalue()


def break_header_case(data, part):
    """Spell one entry's name differently in its local header."""
    raw = bytearray(data)
    needle = part.encode()
    with zipfile.ZipFile(BytesIO(data)) as z:
        info = z.getinfo(part)
    offset = info.header_offset
    (namelen,) = struct.unpack("<H", raw[offset + 26 : offset + 28])
    start = offset + 30
    assert bytes(raw[start : start + namelen]) == needle
    raw[start : start + namelen] = part.upper().encode()[:namelen]
    return bytes(raw)


@pytest.fixture
def sample_bytes():
    return SAMPLES["alpha"].read_bytes()


# --- a reference whose capitalisation does not match the stored part --------


def test_a_miscased_reference_is_repointed(tmp_path, sample_bytes):
    data = rebuild(
        BytesIO(sample_bytes),
        repoint={"docProps/core.xml": "docProps/CORE.xml"},
    )
    path = tmp_path / "miscased.docx"
    path.write_bytes(data)

    with pytest.raises(Exception):
        Document(str(path))

    repaired, notes = repair(path)
    assert any("different capitalisation" in n for n in notes)
    assert Document(BytesIO(repaired)) is not None


def test_the_parser_recovers_such_a_document(tmp_path, sample_bytes):
    path = tmp_path / "recovered.docx"
    path.write_bytes(
        rebuild(BytesIO(sample_bytes), repoint={"docProps/core.xml": "docProps/CORE.xml"})
    )

    doc = parse_document(path)
    assert doc.ok, doc.error
    from conftest import EXPECTED

    assert doc.metrics["matters"].value == EXPECTED["alpha"]["matters"]
    assert doc.metrics["active_clients"].value == EXPECTED["alpha"]["active_clients"]
    assert any("had to be repaired" in w for w in doc.warnings)


# --- an entry named differently in the index and in its own header ---------


def test_a_header_name_mismatch_is_tolerated(tmp_path, sample_bytes):
    path = tmp_path / "header.docx"
    path.write_bytes(break_header_case(sample_bytes, "docProps/core.xml"))

    with pytest.raises(zipfile.BadZipFile):
        with zipfile.ZipFile(path) as z:
            z.read("docProps/core.xml")

    repaired, notes = repair(path)
    assert any("archive's index" in n for n in notes)
    with zipfile.ZipFile(BytesIO(repaired)) as z:
        assert z.read("docProps/core.xml")


def test_the_parser_recovers_a_header_mismatch(tmp_path, sample_bytes):
    path = tmp_path / "header2.docx"
    path.write_bytes(break_header_case(sample_bytes, "docProps/core.xml"))

    doc = parse_document(path)
    assert doc.ok, doc.error
    from conftest import EXPECTED

    assert doc.metrics["matters"].value == EXPECTED["alpha"]["matters"]
    assert any("had to be repaired" in w for w in doc.warnings)


# --- and repair is never used on a file that opens normally ----------------


def test_an_intact_document_is_not_reported_as_repaired():
    doc = parse_document(SAMPLES["alpha"])
    assert doc.ok
    assert not any("had to be repaired" in w for w in doc.warnings)


def test_a_file_that_is_not_a_document_still_fails_clearly(tmp_path):
    path = tmp_path / "nonsense.docx"
    path.write_bytes(b"this is not a docx at all")

    doc = parse_document(path)
    assert not doc.ok
    assert "Could not open the document" in doc.error
    assert all(m.is_unparsed for m in doc.metrics.values())


def test_repair_leaves_the_content_alone(tmp_path, sample_bytes):
    """Only part names change; the document text must be identical."""
    path = tmp_path / "compare.docx"
    path.write_bytes(
        rebuild(BytesIO(sample_bytes), repoint={"docProps/core.xml": "docProps/CORE.xml"})
    )
    repaired, _ = repair(path)

    original = Document(BytesIO(sample_bytes))
    fixed = Document(BytesIO(repaired))
    assert [p.text for p in original.paragraphs] == [p.text for p in fixed.paragraphs]
    assert len(original.tables) == len(fixed.tables)
