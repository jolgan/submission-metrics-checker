"""Reading the new-client column, including answers firms qualify in prose.

A submission that answered "Yes (for this work type)" three times was counted
as 6 new clients instead of 9: the qualified answers fell through to
"unrecognised" and were silently counted as No, with nothing on screen to say
so. Both halves of that are covered here.
"""

import pytest
from docx import Document

from conftest import SAMPLES
from parsing import parse_document

BARE = ["Yes", "yes", "YES", "Y", "y"]
QUALIFIED_YES = [
    "Yes (for this work type)",
    "Yes - new this year",
    "Y (partially)",
    "Yes, first instruction in 2026",
    "YES (new)",
]
QUALIFIED_NO = ["No (existing client)", "N - client for 10 years", "No, since 2015"]
NOT_ANSWERS = ["N/A", "n / a", "N.A.", "Not applicable", "Unknown", "TBC", "?"]


def make_document(path, answers):
    """A minimal submission with one client per answer."""
    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")

    table = document.add_table(rows=len(answers) + 1, cols=2)
    table.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    table.rows[0].cells[1].text = "New client (yes/no)"
    for i, answer in enumerate(answers, start=1):
        table.rows[i].cells[0].text = f"Client {i} Ltd"
        table.rows[i].cells[1].text = answer

    document.save(path)
    return path


def count(tmp_path, answers, name="doc.docx"):
    doc = parse_document(make_document(tmp_path / name, answers))
    return doc.metrics["new_clients"], doc.metrics["active_clients"]


# --- the reported bug -------------------------------------------------------


def test_a_qualified_yes_counts_as_a_new_client(tmp_path):
    """The reported case: 9 yeses, three of them qualified."""
    answers = ["Yes"] * 6 + ["Yes (for this work type)"] * 3
    new, active = count(tmp_path, answers)
    assert active.value == 9
    assert new.value == 9, "qualified answers were being counted as No"


@pytest.mark.parametrize("answer", QUALIFIED_YES)
def test_every_shape_of_qualified_yes(tmp_path, answer):
    new, _ = count(tmp_path, [answer], name=f"{abs(hash(answer))}.docx")
    assert new.value == 1


@pytest.mark.parametrize("answer", QUALIFIED_NO)
def test_every_shape_of_qualified_no(tmp_path, answer):
    new, active = count(tmp_path, [answer], name=f"{abs(hash(answer))}.docx")
    assert new.value == 0
    assert active.value == 1, "it is still an active client"


def test_a_qualified_answer_is_reported_not_silently_interpreted(tmp_path):
    new, _ = count(tmp_path, ["Yes", "Yes (for this work type)"])
    assert new.value == 2
    notes = " ".join(new.notes)
    assert "qualified answer" in notes
    assert "Yes (for this work type)" in notes


def test_bare_answers_produce_no_qualified_note(tmp_path):
    new, _ = count(tmp_path, ["Yes", "No", "Y", "N"])
    assert new.value == 2
    assert not any("qualified" in n for n in new.notes)


# --- the second half: an answer that is neither yes nor no ------------------


@pytest.mark.parametrize("answer", NOT_ANSWERS)
def test_an_unreadable_answer_flags_the_count_for_a_manual_check(tmp_path, answer):
    """Counting these as "not new" would be a guess that hides itself."""
    new, active = count(tmp_path, ["Yes", "No", answer], name=f"{abs(hash(answer))}.docx")
    assert new.is_unparsed, f"{answer!r} should flag the column"
    assert active.value == 3, "the client still counts as active"
    assert any(repr(answer) in n for n in new.notes)
    assert any("by hand" in n for n in new.notes)


def test_n_slash_a_is_not_read_as_no(tmp_path):
    """'N/A' begins with an N but does not mean no."""
    new, _ = count(tmp_path, ["N/A"])
    assert new.is_unparsed


def test_blank_answers_are_still_not_new_and_not_flagged(tmp_path):
    new, active = count(tmp_path, ["Yes", "", ""])
    assert new.value == 1
    assert active.value == 3
    assert not new.is_unparsed
    assert any("blank answer" in n for n in new.notes)


def test_new_and_existing_vocabulary_still_works(tmp_path):
    new, _ = count(tmp_path, ["New", "Existing", "New client from March"])
    assert new.value == 2
    assert any("'New'/'Existing'" in n for n in new.notes)


# --- the samples are unaffected --------------------------------------------


@pytest.mark.parametrize("doc_key", sorted(SAMPLES))
def test_sample_counts_are_unchanged(doc_key):
    from conftest import EXPECTED

    doc = parse_document(SAMPLES[doc_key])
    assert doc.metrics["new_clients"].value == EXPECTED[doc_key]["new_clients"]
    assert not doc.metrics["new_clients"].is_unparsed


# --- "N/A" in the client column ---------------------------------------------
#
# Two submissions had a publishable client table holding a single row reading
# "N/A" - the firm saying it has none - which was counted as a client.


def make_client_table(path, names_and_answers):
    from docx import Document

    document = Document()
    document.add_paragraph("Firm Name")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Example Firm LLP"
    document.add_paragraph("Clients: publishable clients")
    table = document.add_table(rows=len(names_and_answers) + 1, cols=2)
    table.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    table.rows[0].cells[1].text = "New client (yes/no)"
    for i, (name, answer) in enumerate(names_and_answers, start=1):
        table.rows[i].cells[0].text = name
        table.rows[i].cells[1].text = answer
    document.save(path)
    return parse_document(path).metrics


@pytest.mark.parametrize("marker", ["N/A", "n/a", "N / A", "N.A.", "Not applicable"])
def test_not_applicable_in_the_client_column_is_not_a_client(tmp_path, marker):
    metrics = make_client_table(
        tmp_path / f"{abs(hash(marker))}.docx", [(marker, "")]
    )
    assert metrics["active_clients"].value == 0
    assert any("N/A" in n for n in metrics["active_clients"].notes)


def test_it_does_not_disturb_the_real_clients(tmp_path):
    metrics = make_client_table(
        tmp_path / "mixed.docx",
        [("N/A", ""), ("A Client Ltd", "Yes"), ("B Client Ltd", "No")],
    )
    assert metrics["active_clients"].value == 2
    assert metrics["new_clients"].value == 1


def test_a_client_whose_name_merely_contains_na_still_counts(tmp_path):
    metrics = make_client_table(
        tmp_path / "na.docx",
        [("National Grid", "No"), ("Nabarro Holdings", "No"), ("NatWest", "Yes")],
    )
    assert metrics["active_clients"].value == 3


def test_the_na_rule_fires_only_where_there_is_an_na_row():
    from conftest import EXPECTED

    for key in SAMPLES:
        doc = parse_document(SAMPLES[key])
        assert doc.metrics["active_clients"].value == EXPECTED[key]["active_clients"]

    gamma = parse_document(SAMPLES["gamma"]).metrics["active_clients"]
    alpha = parse_document(SAMPLES["alpha"]).metrics["active_clients"]
    assert any("N/A" in n for n in gamma.notes), "gamma has an N/A client row"
    assert not any("N/A" in n for n in alpha.notes), "alpha does not"
