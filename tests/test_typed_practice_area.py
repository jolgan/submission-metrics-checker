"""The 'OR type in this box' fallback, built on synthetic documents.

The four samples all used the dropdown, so this path needs documents made for
it. Each one is a minimal form: the practice-area prompts plus enough of a
client table for the document to be recognised as a submission.
"""

import pytest
from docx import Document

from conftest import POWERBI, PORTAL, SAMPLES
from matching import build_results, load_portal, load_powerbi
from parsing import parse_document

PROMPT = "OR If you have an earlier version of Word, type in this box: ►"


def make_document(path, typed=None, typed_on_next_line=None, extra_paragraphs=()):
    document = Document()
    document.add_paragraph("Firm Name")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Example Firm LLP"

    document.add_paragraph("Practice Area")
    document.add_paragraph("EITHER select Practice Area from this drop-own list ►")
    document.add_paragraph(PROMPT + (f" {typed}" if typed else ""))
    if typed_on_next_line:
        document.add_paragraph(typed_on_next_line)
    for text in extra_paragraphs:
        document.add_paragraph(text)

    clients = document.add_table(rows=3, cols=2)
    clients.rows[0].cells[0].text = "Active key clients (over the last 12 months)"
    clients.rows[0].cells[1].text = "New client (yes/no)"
    clients.rows[1].cells[0].text = "A Client Ltd"
    clients.rows[1].cells[1].text = "Yes"
    clients.rows[2].cells[0].text = "B Client Ltd"
    clients.rows[2].cells[1].text = "No"

    document.save(path)
    return path


def test_typed_on_the_prompt_line_is_read(tmp_path):
    path = make_document(tmp_path / "a.docx", typed="Banking and finance")
    doc = parse_document(path)
    assert doc.typed_practice_area == "Banking and finance"
    assert doc.practice_area == "Banking and finance"
    assert doc.concatenated_practice_area == "Banking and finance"


def test_typed_on_the_following_line_is_read(tmp_path):
    path = make_document(tmp_path / "b.docx", typed_on_next_line="Commercial litigation")
    doc = parse_document(path)
    assert doc.typed_practice_area == "Commercial litigation"


def test_full_dropdown_text_typed_out_is_split(tmp_path):
    path = make_document(
        tmp_path / "c.docx", typed="South West - Finance - Banking and finance"
    )
    doc = parse_document(path)
    assert doc.region == "South West"
    assert doc.practice_group == "Finance"
    assert doc.practice_area == "Banking and finance"
    assert doc.concatenated_practice_area == "Finance > Banking and finance"


def test_form_boilerplate_is_not_mistaken_for_a_typed_answer(tmp_path):
    path = make_document(
        tmp_path / "d.docx",
        extra_paragraphs=[
            "Choose ONE from list of Practice Areas on last page of this document",
            "Contact details to arrange interviews",
        ],
    )
    doc = parse_document(path)
    assert doc.typed_practice_area is None
    assert doc.practice_area is None


def test_an_empty_box_yields_nothing_rather_than_a_guess(tmp_path):
    path = make_document(tmp_path / "e.docx")
    doc = parse_document(path)
    assert doc.typed_practice_area is None
    assert doc.practice_area is None
    assert any("dropdown" in w for w in doc.warnings)


def test_overlong_text_is_rejected(tmp_path):
    path = make_document(tmp_path / "f.docx", typed="x" * 400)
    doc = parse_document(path)
    assert doc.typed_practice_area is None


def test_using_the_typed_box_is_reported_in_the_warnings(tmp_path):
    path = make_document(tmp_path / "g.docx", typed="Corporate tax")
    doc = parse_document(path)
    assert any("typed box" in w for w in doc.warnings)


def test_dropdown_wins_when_both_are_present():
    """The samples all use the dropdown; the typed box must not override it."""
    from conftest import DETAILS

    doc = parse_document(SAMPLES["alpha"])
    assert doc.dropdown_raw
    assert doc.concatenated_practice_area == DETAILS["alpha"]["practice_area"]


# --- interaction with the exports ------------------------------------------


@pytest.fixture(scope="module")
def exports():
    return load_powerbi(POWERBI), load_portal(PORTAL)


def test_typed_value_disagreeing_with_the_portal_is_reported(exports, tmp_path):
    powerbi, portal = exports
    doc = parse_document(SAMPLES["alpha"])
    doc.dropdown_raw = None
    doc.typed_practice_area = "Something the firm typed"
    rows = build_results([doc], powerbi, portal)
    from conftest import DETAILS

    assert rows[0].practice_area == DETAILS["alpha"]["practice_area"]  # portal wins
    assert any("typed box reads" in i for i in rows[0].issues)


def test_typed_value_agreeing_with_the_portal_is_not_flagged(exports):
    powerbi, portal = exports
    doc = parse_document(SAMPLES["alpha"])
    doc.dropdown_raw = None
    doc.typed_practice_area = "Commercial property: Thames Valley"
    rows = build_results([doc], powerbi, portal)
    assert not any("typed box reads" in i for i in rows[0].issues)


def test_typed_value_fills_the_column_when_the_portal_has_no_row(exports, tmp_path):
    """A document missing from the portal export can still be matched."""
    from matching import PortalExport

    powerbi, _ = exports
    doc = parse_document(SAMPLES["gamma"])
    doc.dropdown_raw = None
    doc.region = "London"
    doc.typed_practice_area = "Finance - Banking and finance"
    doc.practice_group = "Finance"
    doc.practice_area = "Banking and finance"
    rows = build_results([doc], powerbi, PortalExport())
    assert rows[0].practice_area == "Finance > Banking and finance"
    assert rows[0].status == "matched"
